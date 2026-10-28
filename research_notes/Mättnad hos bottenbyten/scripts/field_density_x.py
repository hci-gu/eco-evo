"""Absolute prey and consumer densities on the two trawling gradients (Kattegat 2013, Irish Sea
2011; Dryad doi:10.5061/dryad.04dv4 and doi:10.5061/dryad.g3qh8) and the saturation index the
engine's benthic pairs would have at those prey levels.

Run: <python with xlrd, openpyxl> -I scripts/field_density_x.py <data dir> <out csv>

Swept area (Hinz et al. 2017 methods [F]): Irish Sea net opening 16 m x distance towed (file);
Kattegat net opening 25 m x 30 min at 3 kn = 2.78 km (tow length not in the files: J, the Irish
Sea protocol). Catchability q = 1 -> P is a LOWER bound on fish density (flatfish q in otter
trawls ~0.2-0.6, J).
Prey: grab wet mass (5 x 0.1 m2), excluding large infauna (Arctica, Spatangus, Brissopsis,
'ECHINODERM', any specimen > 5 g; Hinz 2017 excluded large infauna as not eaten).
Engine numbers (fgconfig/fg_library.yaml, 223.26): benthivores -> infauna a 0.036, h 99.3, refuge
exposed_share 0.2; x = a h (0.2 B). The engine's a is per tonne of the WHOLE benthivore group; it is
used here only to show where field prey levels sit on the engine's curve (diagnostic).
"""
import collections
import csv
import math
import os
import sys

import openpyxl
import xlrd

D = sys.argv[1]
A_INF, H_INF, PHI = 0.036, 99.3, 0.2
LARGE = ('ARCTICA', 'SPATANGUS', 'ECHINODERM', 'BRISSOPSIS', 'BOLOCERA')
COARSE_K = {'K', 'O', 'Q', 'S'}
SWEPT_IS = {'B': 7.3, 'C': 5.8, 'D': 8.8, 'E': 10.5, 'G': 9.4, 'I': 7.8, 'L': 7.7, 'O': 11.9,
            'R': 3.6, 'S': 6.0, 'T': 2.7}
SWEPT_K = {'A': 0.4, 'B': 1.5, 'C': 0.8, 'D': 1.5, 'E': 0.2, 'F': 2.6, 'G': 4.6, 'H': 4.6,
           'I': 4.5, 'J': 7.9, 'L': 5.7, 'M': 0.4, 'N': 4.5, 'P': 5.4, 'R': 7.4}
CONS_FISH = ('pleuronectes platessa', 'limanda limanda', 'hippoglossoides platessoides')
rows_out = []

# ---------------- Kattegat ----------------
kd = os.path.join(D, 'hiddink2016', 'dryad_04dv4')
grabs = collections.defaultdict(set)
prey = collections.Counter()
with open(os.path.join(kd, 'Infauna Kattegat Biomass all stations!!.csv'), encoding='latin-1') as fh:
    for r in csv.DictReader(fh):
        st = r['STATION'].strip()
        if not st:
            continue
        grabs[st].add(r['GRAB'])
        try:
            w = float(r['WEIGHT (G)'] or 0)
            n = float(r['NO.'] or 1) or 1
        except ValueError:
            continue
        if any(k in r['SPECIES'].upper() for k in LARGE) or w / n > 5:
            continue
        prey[st] += w
sh = xlrd.open_workbook(os.path.join(kd, 'Kattegat otter trawl catches.xls')).sheet_by_index(0)
hauls = collections.defaultdict(set)
fishw = collections.Counter()
nephw = collections.Counter()
codw = collections.Counter()
for i in range(1, sh.nrows):
    v = sh.row_values(i)
    st, sp = v[1].strip(), v[3].strip().lower()
    hauls[st].add(v[2])
    try:
        w = float(v[5]) / 1e6  # g -> t
    except (TypeError, ValueError):
        continue
    if sp in CONS_FISH:
        fishw[st] += w
    elif sp.startswith('nephrops'):
        nephw[st] += w
    elif sp.startswith('gadus'):
        codw[st] += w
SA_K = 25e-3 * 2.78  # km2 per haul (J)
for st in sorted(SWEPT_K):
    if st not in hauls or st in COARSE_K:
        continue
    B = prey[st] / (0.1 * len(grabs[st]))
    nh = len(hauls[st])
    rows_out.append(('Kattegat', st, SWEPT_K[st], B, fishw[st] / nh / SA_K, nephw[st] / nh / SA_K,
                     codw[st] / nh / SA_K))

# ---------------- Irish Sea ----------------
idir = os.path.join(D, 'hinz2017')
wb = openpyxl.load_workbook(os.path.join(idir, 'Benthic infauna_Irish Sea.xlsx'), read_only=True)
ws = wb['Biomass Ww g 0.1 m2']
it = ws.iter_rows(values_only=True)
hdr = next(it)
bsum = collections.Counter()
ngrab = collections.Counter()
for c in hdr[1:]:
    if c:
        ngrab[str(c)[0]] += 1
for r in it:
    name = str(r[0] or '').upper()
    if not name or any(k in name for k in LARGE):
        continue
    for c, v in zip(hdr[1:], r[1:]):
        if c and isinstance(v, (int, float)):
            bsum[str(c)[0]] += v if v <= 5 else 0
wb2 = openpyxl.load_workbook(os.path.join(idir, 'Irish Sea otter trawl and grab positions.xlsx'),
                             read_only=True)
dist = collections.defaultdict(float)
ntow = collections.Counter()
for i, r in enumerate(wb2['Positions otter trawl tows'].iter_rows(values_only=True)):
    if i == 0 or not r[1]:
        continue
    dist[r[1]] += float(r[6])
    ntow[r[1]] += 1
fis = collections.Counter()
nis = collections.Counter()
cis = collections.Counter()
with open(os.path.join(idir, 'Otter trawl catches Irish Sea.csv'), encoding='latin-1') as fh:
    for r in csv.DictReader(fh):
        sp = r['Species'].strip().lower()
        try:
            w = float(r['Weight kg']) / 1e3
        except ValueError:
            continue
        if sp in CONS_FISH:
            fis[r['Station']] += w
        elif sp.startswith('nephrops'):
            nis[r['Station']] += w
        elif sp.startswith('gadus'):
            cis[r['Station']] += w
for st in sorted(SWEPT_IS):
    if st not in dist or not ngrab[st]:
        continue
    sa = 16e-3 * dist[st] / 1e3  # km2 summed over tows
    B = bsum[st] / (0.1 * ngrab[st])
    rows_out.append(('IrishSea', st, SWEPT_IS[st], B, fis[st] / sa, nis[st] / sa, cis[st] / sa))

print('area      st swept/yr  B_prey t/km2  P_flatfish t/km2  P_Nephrops  P_cod   B/P_flat  engine x=a h 0.2B')
with open(sys.argv[2], 'w', newline='') as fh:
    w = csv.writer(fh)
    w.writerow(['area', 'station', 'swept_per_yr', 'B_prey_t_km2', 'P_flatfish_t_km2_q1',
                'P_nephrops_t_km2_q1', 'P_cod_t_km2_q1', 'B_over_Pflat', 'engine_x_infauna'])
    for area, st, sw, B, pf, pn, pc in rows_out:
        x = A_INF * H_INF * PHI * B
        print(f'{area:9s} {st}  {sw:6.1f}  {B:10.1f}  {pf:14.3f}  {pn:10.3f}  {pc:6.3f}  {B / max(pf, 1e-9):8.0f}  {x:8.1f}')
        w.writerow([area, st, sw, f'{B:.2f}', f'{pf:.4f}', f'{pn:.4f}', f'{pc:.4f}',
                    f'{B / max(pf, 1e-9):.1f}', f'{x:.2f}'])


def gm(v):
    return math.exp(sum(math.log(max(a, 1e-6)) for a in v) / len(v))


for area in ('Kattegat', 'IrishSea'):
    rs = [r for r in rows_out if r[0] == area]
    lo = [r for r in rs if r[2] < (3 if area == 'Kattegat' else 6)]
    hi = [r for r in rs if r[2] >= (4.5 if area == 'Kattegat' else 8)]
    print(f'{area}: B_prey gm {gm([r[3] for r in rs]):.1f} t/km2 (range {min(r[3] for r in rs):.1f}-'
          f'{max(r[3] for r in rs):.1f}); flatfish gm {gm([r[4] for r in rs]):.3f} t/km2 at q=1; '
          f'low-trawl B {gm([r[3] for r in lo]):.1f} vs high {gm([r[3] for r in hi]):.1f}; '
          f'flatfish {gm([r[4] for r in lo]):.3f} vs {gm([r[4] for r in hi]):.3f}')
    print(f'   engine x on infauna at these B: {A_INF * H_INF * PHI * min(r[3] for r in rs):.1f}-'
          f'{A_INF * H_INF * PHI * max(r[3] for r in rs):.1f}; '
          f'Holling-II intake ratio high/low trawl at x_low = engine: '
          f'{(gm([r[3] for r in hi]) / gm([r[3] for r in lo])) * (1 + A_INF * H_INF * PHI * gm([r[3] for r in lo])) / (1 + A_INF * H_INF * PHI * gm([r[3] for r in hi])):.3f}; '
          f'at x_low = 2: {(gm([r[3] for r in hi]) / gm([r[3] for r in lo])) * 3 / (1 + 2 * gm([r[3] for r in hi]) / gm([r[3] for r in lo])):.3f}')
